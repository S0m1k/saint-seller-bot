"""Статистика по подписчикам и выгрузка в Excel."""
from datetime import datetime, timedelta, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from core.config import settings
from core.models import Order, User


def _tz() -> ZoneInfo:
    try:
        return ZoneInfo(settings.timezone)
    except Exception:  # noqa: BLE001
        return ZoneInfo("UTC")


def _now_utc() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def local_day_bounds_utc(days_ago: int = 0) -> tuple[datetime, datetime]:
    """Границы календарного дня в локальном TZ, приведённые к naive-UTC.

    days_ago=0 — сегодня, 1 — вчера. created_at в БД хранится как naive-UTC
    (SQLite CURRENT_TIMESTAMP), поэтому сравниваем в UTC.
    """
    tz = _tz()
    now_local = datetime.now(tz)
    day_local = (now_local - timedelta(days=days_ago)).replace(
        hour=0, minute=0, second=0, microsecond=0
    )
    start_utc = day_local.astimezone(timezone.utc).replace(tzinfo=None)
    end_utc = (day_local + timedelta(days=1)).astimezone(timezone.utc).replace(tzinfo=None)
    return start_utc, end_utc


async def _count_since(session: AsyncSession, start_utc: datetime, end_utc: datetime | None = None) -> int:
    stmt = select(func.count(User.telegram_id)).where(User.created_at >= start_utc)
    if end_utc is not None:
        stmt = stmt.where(User.created_at < end_utc)
    return (await session.execute(stmt)).scalar_one()


async def collect_stats(session: AsyncSession) -> dict:
    """Сводные показатели по подписчикам и заказам."""
    total = (await session.execute(select(func.count(User.telegram_id)))).scalar_one()
    active = (await session.execute(
        select(func.count(User.telegram_id)).where(User.is_active.is_(True))
    )).scalar_one()

    today_start, today_end = local_day_bounds_utc(0)
    yest_start, yest_end = local_day_bounds_utc(1)
    week_start = _now_utc() - timedelta(days=7)
    month_start = _now_utc() - timedelta(days=30)

    new_today = await _count_since(session, today_start, today_end)
    new_yesterday = await _count_since(session, yest_start, yest_end)
    new_week = await _count_since(session, week_start)
    new_month = await _count_since(session, month_start)

    orders_total = (await session.execute(select(func.count(Order.id)))).scalar_one()

    return {
        "total": total,
        "active": active,
        "new_today": new_today,
        "new_yesterday": new_yesterday,
        "new_week": new_week,
        "new_month": new_month,
        "orders_total": orders_total,
    }


def format_stats(stats: dict, title: str = "📊 <b>Статистика</b>") -> str:
    return (
        f"{title}\n\n"
        f"👥 Всего подписчиков: <b>{stats['total']}</b>\n"
        f"✅ Активных: <b>{stats['active']}</b>\n"
        f"🧾 Заказов всего: <b>{stats['orders_total']}</b>\n\n"
        f"<b>Новые подписчики:</b>\n"
        f"• Сегодня: <b>{stats['new_today']}</b>\n"
        f"• Вчера: <b>{stats['new_yesterday']}</b>\n"
        f"• За 7 дней: <b>{stats['new_week']}</b>\n"
        f"• За 30 дней: <b>{stats['new_month']}</b>"
    )


async def fetch_users(
    session: AsyncSession,
    start_utc: datetime | None = None,
    end_utc: datetime | None = None,
) -> list[User]:
    stmt = select(User).order_by(User.created_at.desc())
    if start_utc is not None:
        stmt = stmt.where(User.created_at >= start_utc)
    if end_utc is not None:
        stmt = stmt.where(User.created_at < end_utc)
    return list((await session.execute(stmt)).scalars().all())


def users_to_xlsx(users: list[User]) -> bytes:
    """Формирует xlsx-файл со списком подписчиков (даты — в локальном TZ)."""
    tz = _tz()
    wb = Workbook()
    ws = wb.active
    ws.title = "Подписчики"

    headers = ["Telegram ID", "Username", "Имя", "Фамилия", "Телефон", "Активен", "Дата регистрации"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for u in users:
        created_local = ""
        if u.created_at:
            created_local = (
                u.created_at.replace(tzinfo=timezone.utc)
                .astimezone(tz)
                .strftime("%d.%m.%Y %H:%M")
            )
        ws.append([
            u.telegram_id,
            f"@{u.username}" if u.username else "",
            u.first_name or "",
            u.last_name or "",
            u.phone or "",
            "да" if u.is_active else "нет",
            created_local,
        ])

    widths = [16, 20, 18, 18, 18, 10, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
